from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from database import get_session
from models import Account, Expense, Goal, Investment, Todo

router = APIRouter(prefix="/export", tags=["Export"])

HEADER_FILL = PatternFill("solid", start_color="3B82F6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '"₹"#,##0.00'
PERCENT = "0.0%"
DATE_FMT = "DD-MMM-YYYY"


def add_sheet(wb, title, headers, rows, money_cols=(), percent_cols=(), date_cols=()):
    """Create one formatted sheet. Column numbers start at 1 (A=1, B=2, ...)."""
    ws = wb.create_sheet(title)
    ws.append(headers)

    # Style the header row
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"  # header stays visible when scrolling

    # Number formats
    for r in range(2, ws.max_row + 1):
        for c in money_cols:
            ws.cell(r, c).number_format = MONEY
        for c in percent_cols:
            ws.cell(r, c).number_format = PERCENT
        for c in date_cols:
            ws.cell(r, c).number_format = DATE_FMT

    # Auto-fit column widths (formulas get a fixed width)
    for i, header in enumerate(headers, start=1):
        widths = [len(str(header))]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, i).value
            widths.append(14 if isinstance(v, str) and v.startswith("=") else len(str(v or "")))
        ws.column_dimensions[get_column_letter(i)].width = min(max(widths) + 3, 40)

    return ws


@router.get("/excel")
def export_excel(session: Session = Depends(get_session)):
    expenses = session.exec(select(Expense).order_by(Expense.date.desc())).all()
    goals = session.exec(select(Goal)).all()
    accounts = session.exec(select(Account)).all()
    todos = session.exec(select(Todo)).all()
    investments = session.exec(select(Investment)).all()

    wb = Workbook()

    # ---------- Summary sheet (first tab) ----------
    today = date.today()
    cash = sum(a.balance for a in accounts)
    invested = sum(i.units * i.buy_price for i in investments)
    inv_value = sum(i.units * i.current_value for i in investments)
    month_spend = sum(
        e.amount for e in expenses if e.date.year == today.year and e.date.month == today.month
    )

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Life OS Report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated on {today.strftime('%d-%b-%Y')}"

    summary_rows = [
        ("Cash in accounts", cash),
        ("Amount invested", invested),
        ("Investments current value", inv_value),
        ("Investment gain", inv_value - invested),
        ("Net worth", cash + inv_value),
        (f"Spending in {today.strftime('%B %Y')}", month_spend),
    ]
    for n, (label, value) in enumerate(summary_rows, start=4):
        ws.cell(n, 1, label).font = Font(bold=True)
        ws.cell(n, 2, round(value, 2)).number_format = MONEY
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    # ---------- Expenses ----------
    add_sheet(
        wb, "Expenses",
        ["Date", "Category", "Amount", "Notes"],
        [[e.date, e.category, e.amount, e.notes] for e in expenses],
        money_cols=[3], date_cols=[1],
    )

    # ---------- Goals (with formulas) ----------
    goal_rows = []
    for n, g in enumerate(goals, start=2):  # row 1 is the header
        goal_rows.append([
            g.name, g.target_amount, g.current_amount,
            f"=IF(B{n}=0,0,C{n}/B{n})",   # % complete
            f"=MAX(B{n}-C{n},0)",         # remaining
            g.target_date, g.category,
        ])
    add_sheet(
        wb, "Goals",
        ["Goal", "Target", "Saved so far", "% complete", "Remaining", "Target date", "Category"],
        goal_rows,
        money_cols=[2, 3, 5], percent_cols=[4], date_cols=[6],
    )

    # ---------- Accounts ----------
    add_sheet(
        wb, "Accounts",
        ["Account", "Type", "Balance"],
        [[a.name, a.account_type, a.balance] for a in accounts],
        money_cols=[3],
    )

    # ---------- Investments (with formulas) ----------
    inv_rows = []
    for n, i in enumerate(investments, start=2):
        inv_rows.append([
            i.name, i.investment_type, i.units, i.buy_price, i.current_value,
            f"=C{n}*D{n}",              # invested
            f"=C{n}*E{n}",              # current value
            f"=G{n}-F{n}",              # gain
            f"=IF(F{n}=0,0,H{n}/F{n})", # gain %
        ])
    add_sheet(
        wb, "Investments",
        ["Name", "Type", "Units", "Buy price/unit", "Current price/unit",
         "Invested", "Current value", "Gain", "Gain %"],
        inv_rows,
        money_cols=[4, 5, 6, 7, 8], percent_cols=[9],
    )

    # ---------- Todos ----------
    add_sheet(
        wb, "Todos",
        ["Task", "Due date", "Status", "Priority"],
        [[t.title, t.due_date, t.status, t.priority] for t in todos],
        date_cols=[2],
    )

    # Save into memory (not onto disk) and send as a download
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"life_os_report_{today.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )